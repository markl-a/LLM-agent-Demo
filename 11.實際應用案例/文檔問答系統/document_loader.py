#!/usr/bin/env python3
"""
文檔加載器模組

提供高級文檔加載功能，支持多種文件格式和自定義處理。
"""

from pathlib import Path
from typing import List, Optional, Dict, Any
import logging

logger = logging.getLogger(__name__)

try:
    from llama_index.core import Document
    from llama_index.core.readers.base import BaseReader
except ImportError:
    logger.error("請安裝 llama-index: pip install llama-index")
    raise


class EnhancedDocumentLoader:
    """增強型文檔加載器"""

    # 支援的文件類型
    SUPPORTED_EXTENSIONS = {
        ".txt": "text",
        ".md": "markdown",
        ".pdf": "pdf",
        ".docx": "word",
        ".xlsx": "excel",
        ".csv": "csv",
        ".json": "json",
        ".html": "html",
        ".htm": "html",
    }

    def __init__(self):
        """初始化文檔加載器"""
        self.documents: List[Document] = []

    def load_directory(
        self,
        directory: str,
        recursive: bool = True,
        required_exts: Optional[List[str]] = None,
        exclude_patterns: Optional[List[str]] = None,
    ) -> List[Document]:
        """
        載入目錄中的所有文檔

        Args:
            directory: 目錄路徑
            recursive: 是否遞迴子目錄
            required_exts: 需要的文件擴展名列表
            exclude_patterns: 排除的文件模式

        Returns:
            文檔列表
        """
        from llama_index.core import SimpleDirectoryReader

        dir_path = Path(directory)
        if not dir_path.exists():
            raise FileNotFoundError(f"目錄不存在: {directory}")

        if not dir_path.is_dir():
            raise ValueError(f"路徑不是目錄: {directory}")

        # 設置默認擴展名
        if required_exts is None:
            required_exts = list(self.SUPPORTED_EXTENSIONS.keys())

        # 設置默認排除模式
        if exclude_patterns is None:
            exclude_patterns = [
                ".*",  # 隱藏文件
                "__pycache__",
                "*.pyc",
                "node_modules",
                ".git",
            ]

        logger.info(f"載入目錄: {directory}")
        logger.info(f"遞迴: {recursive}")
        logger.info(f"支援的擴展名: {required_exts}")

        try:
            reader = SimpleDirectoryReader(
                directory,
                recursive=recursive,
                required_exts=required_exts,
                exclude=exclude_patterns,
            )

            documents = reader.load_data()
            self.documents.extend(documents)

            logger.info(f"成功載入 {len(documents)} 個文檔")

            return documents

        except Exception as e:
            logger.error(f"載入目錄失敗: {e}")
            raise

    def load_file(self, file_path: str, extra_metadata: Optional[Dict[str, Any]] = None) -> Document:
        """
        載入單個文件

        Args:
            file_path: 文件路徑
            extra_metadata: 額外的元數據

        Returns:
            Document 實例
        """
        file = Path(file_path)

        if not file.exists():
            raise FileNotFoundError(f"文件不存在: {file_path}")

        if not file.is_file():
            raise ValueError(f"路徑不是文件: {file_path}")

        # 檢查文件類型
        ext = file.suffix.lower()
        if ext not in self.SUPPORTED_EXTENSIONS:
            logger.warning(f"不支援的文件類型: {ext}")

        logger.info(f"載入文件: {file_path}")

        try:
            # 讀取文件內容
            with open(file_path, "r", encoding="utf-8") as f:
                content = f.read()

            # 創建元數據
            metadata = {
                "file_name": file.name,
                "file_path": str(file.absolute()),
                "file_type": self.SUPPORTED_EXTENSIONS.get(ext, "unknown"),
                "file_size": file.stat().st_size,
            }

            # 添加額外元數據
            if extra_metadata:
                metadata.update(extra_metadata)

            # 創建文檔
            document = Document(text=content, metadata=metadata)

            self.documents.append(document)

            logger.info(f"成功載入文件: {file.name}")

            return document

        except UnicodeDecodeError:
            logger.error(f"文件編碼錯誤: {file_path}")
            logger.info("嘗試使用 PDF/Word 專用加載器")
            return self._load_binary_file(file_path, extra_metadata)

        except Exception as e:
            logger.error(f"載入文件失敗: {e}")
            raise

    def _load_binary_file(
        self, file_path: str, extra_metadata: Optional[Dict[str, Any]] = None
    ) -> Document:
        """
        載入二進制文件（PDF, Word 等）

        Args:
            file_path: 文件路徑
            extra_metadata: 額外的元數據

        Returns:
            Document 實例
        """
        file = Path(file_path)
        ext = file.suffix.lower()

        try:
            if ext == ".pdf":
                return self._load_pdf(file_path, extra_metadata)
            elif ext == ".docx":
                return self._load_word(file_path, extra_metadata)
            elif ext == ".xlsx":
                return self._load_excel(file_path, extra_metadata)
            else:
                raise ValueError(f"不支援的二進制文件類型: {ext}")

        except ImportError as e:
            logger.error(f"缺少必要的庫: {e}")
            logger.info("請安裝: pip install pypdf python-docx openpyxl")
            raise

    def _load_pdf(self, file_path: str, extra_metadata: Optional[Dict[str, Any]] = None) -> Document:
        """載入 PDF 文件"""
        from pypdf import PdfReader

        reader = PdfReader(file_path)
        text = ""

        for page_num, page in enumerate(reader.pages, 1):
            text += f"\n--- 第 {page_num} 頁 ---\n"
            text += page.extract_text()

        metadata = {
            "file_name": Path(file_path).name,
            "file_path": str(Path(file_path).absolute()),
            "file_type": "pdf",
            "total_pages": len(reader.pages),
        }

        if extra_metadata:
            metadata.update(extra_metadata)

        return Document(text=text, metadata=metadata)

    def _load_word(self, file_path: str, extra_metadata: Optional[Dict[str, Any]] = None) -> Document:
        """載入 Word 文件"""
        from docx import Document as WordDocument

        doc = WordDocument(file_path)
        text = "\n\n".join([para.text for para in doc.paragraphs if para.text.strip()])

        metadata = {
            "file_name": Path(file_path).name,
            "file_path": str(Path(file_path).absolute()),
            "file_type": "word",
            "total_paragraphs": len(doc.paragraphs),
        }

        if extra_metadata:
            metadata.update(extra_metadata)

        return Document(text=text, metadata=metadata)

    def _load_excel(self, file_path: str, extra_metadata: Optional[Dict[str, Any]] = None) -> Document:
        """載入 Excel 文件"""
        from openpyxl import load_workbook

        workbook = load_workbook(file_path, data_only=True)
        text = ""

        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            text += f"\n\n=== 工作表: {sheet_name} ===\n\n"

            for row in sheet.iter_rows(values_only=True):
                row_text = "\t".join([str(cell) if cell is not None else "" for cell in row])
                text += row_text + "\n"

        metadata = {
            "file_name": Path(file_path).name,
            "file_path": str(Path(file_path).absolute()),
            "file_type": "excel",
            "total_sheets": len(workbook.sheetnames),
            "sheet_names": workbook.sheetnames,
        }

        if extra_metadata:
            metadata.update(extra_metadata)

        return Document(text=text, metadata=metadata)

    def get_statistics(self) -> Dict[str, Any]:
        """
        獲取載入的文檔統計信息

        Returns:
            統計信息字典
        """
        if not self.documents:
            return {"total_documents": 0}

        # 按文件類型分組
        type_counts = {}
        total_size = 0

        for doc in self.documents:
            file_type = doc.metadata.get("file_type", "unknown")
            type_counts[file_type] = type_counts.get(file_type, 0) + 1

            file_size = doc.metadata.get("file_size", 0)
            total_size += file_size

        return {
            "total_documents": len(self.documents),
            "total_size_bytes": total_size,
            "total_size_mb": round(total_size / (1024 * 1024), 2),
            "types": type_counts,
        }

    def clear(self):
        """清除已載入的文檔"""
        self.documents.clear()
        logger.info("已清除所有文檔")


def main():
    """測試函數"""
    loader = EnhancedDocumentLoader()

    # 測試載入目錄
    try:
        docs = loader.load_directory("./data", recursive=True)
        print(f"\n載入了 {len(docs)} 個文檔")

        # 顯示統計
        stats = loader.get_statistics()
        print("\n統計信息:")
        for key, value in stats.items():
            print(f"  {key}: {value}")

    except FileNotFoundError:
        print("測試目錄不存在，跳過測試")


if __name__ == "__main__":
    main()
